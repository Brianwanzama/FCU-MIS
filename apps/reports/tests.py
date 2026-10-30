"""
Reports and exports: the figures, the files, and who may download them.
"""
import io
from datetime import timedelta
from decimal import Decimal

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.contributions import services as cs
from apps.contributions.models import Contribution
from apps.cycles import services as cycle_services
from apps.cycles.models import Cycle, _add_months
from apps.loans import services as loan_services
from apps.loans.tests import pdf as pdf_upload
from apps.members.models import Member
from apps.unittrust import services as inv_services
from apps.unittrust.models import InvestmentTransaction

from . import excel, services

K = InvestmentTransaction.Kind


class ReportCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        mk = lambda name, role, member=None: User.objects.create_user(
            username=name, password="pw", role=role, member=member
        )
        cls.m1, cls.m2, cls.m3 = Member.objects.order_by("member_code")[:3]
        cls.admin = mk("admin", User.Role.ADMINISTRATOR)
        cls.treasurer = mk("treasurer", User.Role.TREASURER)
        cls.chair = mk("chair", User.Role.CHAIRPERSON)
        cls.secretary = mk("secretary", User.Role.SECRETARY)
        cls.u1 = mk("u1", User.Role.MEMBER, cls.m1)

        first = timezone.localdate().replace(day=1)
        cls.cycle = cycle_services.create_cycle(
            Cycle(cycle_number=1, start_date=_add_months(first, -2), duration_months=6), actor=cls.admin
        )
        cycle_services.activate_cycle(cls.cycle, actor=cls.admin)

    def pay(self, member, amount, kind="MONTHLY", reference=""):
        return cs.record_contribution(
            Contribution(
                member=member, cycle=self.cycle, contribution_type=kind, amount=Decimal(amount),
                period_month=self.cycle.start_date, payment_date=timezone.localdate(),
                payment_method="CASH", reference=reference, recorded_by=self.admin,
            ),
            actor=self.admin,
        )


class ReportFigureTests(ReportCase):
    def test_contribution_figures_and_member_rows(self):
        self.pay(self.m1, "600000")
        self.pay(self.m1, "30000", kind="EMERGENCY")
        self.pay(self.m2, "250000")
        r = services.cycle_report(self.cycle)
        c = r["contributions"]
        self.assertEqual(c["collected"], Decimal("850000"))
        self.assertEqual(c["emergency_collected"], Decimal("30000"))
        self.assertEqual(c["members_met_target"], 1)
        rows = {row["member"].member_code: row for row in r["members"]}
        self.assertTrue(rows[self.m1.member_code]["target_met"])
        self.assertEqual(rows[self.m2.member_code]["deposits_outstanding"], Decimal("350000"))

    def test_loan_section(self):
        self.pay(self.m1, "600000")
        loan_services.upload_undertaking(self.m1, pdf_upload("d.pdf"), self.u1)
        u = self.m1.undertakings.get()
        loan_services.review_undertaking(u, self.chair, verified=True)
        app = loan_services.submit_application(self.m1, Decimal("300000"), 3, pdf_upload("a.pdf"), self.u1)
        loan_services.approve_application(app, self.admin)
        loan = loan_services.disburse_loan(app, self.treasurer, "CASH")
        loan_services.record_repayment(loan, Decimal("110163"), timezone.localdate(), "CASH", self.treasurer)
        s = services.cycle_report(self.cycle)["loans"]
        self.assertEqual(s["count"], 1)
        self.assertEqual(s["disbursed"], Decimal("300000"))
        self.assertEqual(s["repayments_received"], Decimal("110163"))
        self.assertEqual(s["interest_earned"], Decimal("15000"))  # first instalment's interest, paid first
        self.assertEqual(s["outstanding"], loan.outstanding)

    def test_investment_section_and_distribution(self):
        inv = inv_services.create_investment(
            self.treasurer, name="UT", kind="UNIT_TRUST", opened_on=timezone.localdate() - timedelta(days=5)
        )
        inv_services.record_transaction(inv, K.INVEST, Decimal("1000000"), timezone.localdate(), self.treasurer)
        inv_services.record_transaction(inv, K.INCOME, Decimal("50000"), timezone.localdate(), self.treasurer)
        s = services.cycle_report(self.cycle)["investments"]
        self.assertEqual(s["invested_in_cycle"], Decimal("1000000"))
        self.assertEqual(s["income_in_cycle"], Decimal("50000"))
        self.assertEqual(s["market_value"], Decimal("1000000"))

    def test_derived_cash_follows_every_movement_of_money(self):
        self.pay(self.m1, "600000")
        self.pay(self.m1, "30000", kind="EMERGENCY")
        self.assertEqual(services.financial_position()["derived_cash"], Decimal("630000"))
        inv = inv_services.create_investment(
            self.treasurer, name="UT", kind="UNIT_TRUST", opened_on=timezone.localdate() - timedelta(days=5)
        )
        inv_services.record_transaction(inv, K.INVEST, Decimal("400000"), timezone.localdate(), self.treasurer)
        p = services.financial_position()
        self.assertEqual(p["derived_cash"], Decimal("230000"))
        self.assertEqual(p["investments"], Decimal("400000"))
        self.assertEqual(p["members_deposits"], Decimal("600000"))
        self.assertEqual(p["emergency_fund"], Decimal("30000"))

    def test_reversed_payments_are_left_out(self):
        c = self.pay(self.m1, "100000")
        r = cs.request_reversal(c, self.admin, "Duplicate")
        cs.approve_reversal(r, self.chair)
        self.assertEqual(services.cycle_report(self.cycle)["contributions"]["collected"], Decimal("0"))
        self.assertEqual(services.financial_position()["derived_cash"], Decimal("0"))


class ExcelTests(ReportCase):
    def open(self, content):
        return load_workbook(io.BytesIO(content))

    def test_cycle_workbook_has_the_sheets_and_numbers(self):
        self.pay(self.m1, "600000")
        wb = self.open(excel.cycle_workbook(services.cycle_report(self.cycle)))
        self.assertEqual(wb.sheetnames, ["Summary", "Members", "Loans", "Investments"])
        values = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
        self.assertEqual(values["Deposits collected"], 600000)
        member_row = [r for r in wb["Members"].iter_rows(min_row=2, values_only=True) if r[0] == self.m1.member_code][0]
        self.assertEqual(member_row[4], 600000)
        self.assertEqual(member_row[10], "Yes")

    def test_text_that_looks_like_a_formula_stays_text(self):
        self.pay(self.m1, "100000", reference="=HYPERLINK(\"http://evil\",\"x\")")
        self.pay(self.m1, "100000", reference="+1+1")
        wb = self.open(excel.contributions_workbook(Contribution.objects.select_related("member", "cycle", "settles_cycle", "recorded_by")))
        refs = [r[9] for r in wb["Contributions"].iter_rows(min_row=2, values_only=True)]
        self.assertTrue(all(not str(v).startswith(("=", "+")) for v in refs))
        self.assertIn("'=HYPERLINK", " ".join(refs))
        for row in wb["Contributions"].iter_rows(min_row=2):
            for cell in row:
                self.assertNotEqual(cell.data_type, "f", cell.coordinate)

    def test_contributions_workbook_lists_arrears_and_reversals(self):
        c = self.pay(self.m1, "100000")
        r = cs.request_reversal(c, self.admin, "Dup")
        cs.approve_reversal(r, self.chair)
        wb = self.open(excel.contributions_workbook(Contribution.objects.select_related("member", "cycle", "settles_cycle", "recorded_by")))
        row = list(wb["Contributions"].iter_rows(min_row=2, values_only=True))[0]
        self.assertEqual(row[7], 100000)
        self.assertTrue(row[11])  # reversed date shown


class ReportViewTests(ReportCase):
    def login(self, name):
        self.client.logout()
        self.client.login(username=name, password="pw")

    def test_only_officers_who_see_finances_can_open_reports(self):
        urls = [
            "/reports/", f"/reports/cycle/{self.cycle.pk}/", f"/reports/cycle/{self.cycle.pk}/pdf/",
            f"/reports/cycle/{self.cycle.pk}/excel/", "/reports/export/contributions/",
            "/reports/export/loans/", "/reports/export/investments/",
        ]
        for name, code in [("admin", 200), ("treasurer", 200), ("chair", 200), ("secretary", 403), ("u1", 403)]:
            self.login(name)
            for url in urls:
                self.assertEqual(self.client.get(url).status_code, code, (name, url))

    def test_anonymous_is_sent_to_login(self):
        self.client.logout()
        self.assertEqual(self.client.get("/reports/").status_code, 302)

    def test_pdf_and_excel_downloads(self):
        self.pay(self.m1, "600000")
        self.login("treasurer")
        r = self.client.get(f"/reports/cycle/{self.cycle.pk}/pdf/")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("attachment", r["Content-Disposition"])
        r = self.client.get(f"/reports/cycle/{self.cycle.pk}/excel/")
        self.assertTrue(r.content.startswith(b"PK"))  # xlsx is a zip file
        self.assertIn("spreadsheetml", r["Content-Type"])

    def test_html_report_shows_the_figures(self):
        self.pay(self.m1, "600000")
        self.login("chair")
        r = self.client.get(f"/reports/cycle/{self.cycle.pk}/")
        self.assertContains(r, "UGX 600,000")
        self.assertContains(r, "Financial position")

    def test_upcoming_cycles_have_no_report(self):
        upcoming = cycle_services.create_cycle(
            Cycle(cycle_number=2, start_date=_add_months(self.cycle.end_date, 1), duration_months=6), actor=self.admin
        )
        self.login("admin")
        self.assertEqual(self.client.get(f"/reports/cycle/{upcoming.pk}/").status_code, 404)

    def test_contribution_export_can_filter_by_cycle(self):
        self.pay(self.m1, "100000")
        self.login("admin")
        r = self.client.get(f"/reports/export/contributions/?cycle={self.cycle.pk}")
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb["Contributions"].max_row, 2)

    def test_member_statement_excel_follows_statement_access(self):
        self.pay(self.m1, "100000")
        url = f"/members/{self.m1.member_code}/statement.xlsx?cycle={self.cycle.pk}"
        self.login("u1")
        self.assertEqual(self.client.get(url).status_code, 200)
        self.login("treasurer")
        self.assertEqual(self.client.get(url).status_code, 200)
        self.login("secretary")
        self.assertEqual(self.client.get(url).status_code, 403)
        u2 = User.objects.create_user(username="u2", password="pw", role=User.Role.MEMBER, member=self.m2)
        self.client.login(username="u2", password="pw")
        self.assertEqual(self.client.get(url).status_code, 403)


class ChartDataTests(ReportCase):
    def test_monthly_deposits_group_by_the_month_paid_for(self):
        from datetime import date as d

        from .charts import monthly_deposits

        for month, amount in ((self.cycle.start_date, "100000"), (self.cycle.start_date, "50000")):
            cs.record_contribution(
                Contribution(
                    member=self.m1, cycle=self.cycle, contribution_type="MONTHLY", amount=Decimal(amount),
                    period_month=month, payment_date=timezone.localdate(), payment_method="CASH",
                    recorded_by=self.admin,
                ),
                actor=self.admin,
            )
        data = monthly_deposits(self.cycle)
        self.assertEqual(len(data["labels"]), 6)
        self.assertEqual(data["collected"][0], 150000)
        self.assertEqual(sum(data["collected"][1:]), 0)
        n = self.cycle.eligible_member_count
        self.assertEqual(data["target"][0], 100000 * n)
        mine = monthly_deposits(self.cycle, member=self.m1)
        self.assertEqual(mine["target"][0], 100000)
        self.assertEqual(monthly_deposits(self.cycle, member=self.m2)["collected"][0], 0)

    def test_reversed_and_emergency_payments_are_left_out_of_the_monthly_chart(self):
        from .charts import monthly_deposits

        c = self.pay(self.m1, "100000")
        self.pay(self.m1, "30000", kind="EMERGENCY")
        r = cs.request_reversal(c, self.admin, "Dup")
        cs.approve_reversal(r, self.chair)
        self.assertEqual(sum(monthly_deposits(self.cycle)["collected"]), 0)

    def test_member_status_counts_add_up(self):
        from .charts import member_status

        self.pay(self.m1, "600000")
        self.pay(self.m2, "100000")
        rows = services.member_rows(self.cycle)
        data = member_status(rows)
        self.assertEqual(sum(data["values"]), len(rows))
        met, part, none = data["values"]
        self.assertGreaterEqual(met, 1)
        self.assertGreaterEqual(part, 1)
        # every status carries a text label and a symbol, never colour alone
        self.assertTrue(all(len(label) > 2 for label in data["labels"]))

    def test_member_progress_is_sorted_and_links_to_the_member(self):
        from .charts import member_progress

        self.pay(self.m1, "600000")
        self.pay(self.m2, "250000")
        data = member_progress(services.member_rows(self.cycle))
        self.assertEqual(data["paid"][0], 600000)
        self.assertEqual(data["outstanding"][0], 0)
        self.assertEqual(data["paid"][1], 250000)
        self.assertEqual(data["outstanding"][1], 350000)
        self.assertEqual(data["links"][0], f"/members/{self.m1.member_code}/")
        self.assertEqual(data["paid"], sorted(data["paid"], reverse=True))

    def test_allocation_keeps_three_by_name_and_folds_the_rest_into_other(self):
        from .charts import portfolio_allocation

        for i, amount in enumerate(("500000", "400000", "300000", "200000", "100000")):
            inv = inv_services.create_investment(
                self.treasurer, name=f"Fund {i}", kind="UNIT_TRUST",
                opened_on=timezone.localdate() - timedelta(days=5),
            )
            inv_services.record_transaction(inv, K.INVEST, Decimal(amount), timezone.localdate(), self.treasurer)
        data = portfolio_allocation(inv_services.portfolio())
        self.assertEqual(data["labels"], ["Fund 0", "Fund 1", "Fund 2", "Other"])
        self.assertEqual(data["values"], [500000, 400000, 300000, 300000])

    def test_investment_returns_in_percent(self):
        from .charts import investment_returns

        inv = inv_services.create_investment(
            self.treasurer, name="UT", kind="UNIT_TRUST", opened_on=timezone.localdate() - timedelta(days=5)
        )
        inv_services.record_transaction(inv, K.INVEST, Decimal("1000000"), timezone.localdate(), self.treasurer)
        inv_services.record_transaction(inv, K.VALUATION, Decimal("1100000"), timezone.localdate(), self.treasurer)
        data = investment_returns(inv_services.portfolio())
        self.assertEqual(data["labels"], ["UT"])
        self.assertEqual(data["values"], [10.0])


class ChartPageTests(ReportCase):
    def login(self, name):
        self.client.logout()
        self.client.login(username=name, password="pw")

    def test_the_chart_library_ships_with_the_app(self):
        from django.contrib.staticfiles import finders

        self.assertIsNotNone(finders.find("vendor/chart.umd.js"))
        self.assertIsNotNone(finders.find("js/fcu-charts.js"))

    def test_officer_dashboard_has_the_unit_charts(self):
        self.pay(self.m1, "100000")
        self.login("treasurer")
        r = self.client.get("/dashboard/")
        self.assertContains(r, 'id="chart-monthly"')
        self.assertContains(r, 'id="chart-status"')
        self.assertContains(r, "vendor/chart.umd.js")
        self.assertContains(r, 'role="img"')

    def test_member_dashboard_shows_only_their_own_chart(self):
        self.pay(self.m1, "100000")
        self.login("u1")
        r = self.client.get("/dashboard/")
        self.assertContains(r, 'id="chart-mine"')
        self.assertNotContains(r, 'id="chart-status"')
        self.assertNotContains(r, 'id="chart-progress"')

    def test_dashboard_without_an_active_cycle_still_loads(self):
        cycle_services.close_cycle(self.cycle, actor=self.admin)
        self.login("treasurer")
        r = self.client.get("/dashboard/")
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, 'id="chart-monthly"')

    def test_report_page_has_all_four_charts(self):
        self.pay(self.m1, "600000")
        self.login("chair")
        r = self.client.get(f"/reports/cycle/{self.cycle.pk}/")
        for chart in ("chart-monthly", "chart-status", "chart-progress"):
            self.assertContains(r, f'id="{chart}"')
        # no investments yet: a plain message instead of an empty chart
        self.assertNotContains(r, 'id="chart-allocation"')
        self.assertContains(r, "No investments are held yet.")

    def test_member_page_has_the_members_monthly_chart(self):
        self.login("u1")
        r = self.client.get(f"/members/{self.m1.member_code}/")
        self.assertContains(r, 'id="chart-member-monthly"')

    def test_investments_page_charts(self):
        inv = inv_services.create_investment(
            self.treasurer, name="UT", kind="UNIT_TRUST", opened_on=timezone.localdate() - timedelta(days=5)
        )
        inv_services.record_transaction(inv, K.INVEST, Decimal("1000000"), timezone.localdate(), self.treasurer)
        self.login("u1")
        r = self.client.get("/investments/")
        self.assertContains(r, 'id="chart-allocation"')
        self.assertContains(r, 'id="chart-returns"')

    def test_names_cannot_break_out_of_the_embedded_chart_data(self):
        inv = inv_services.create_investment(
            self.treasurer, name="</script><script>alert(1)</script>", kind="UNIT_TRUST",
            opened_on=timezone.localdate() - timedelta(days=5),
        )
        inv_services.record_transaction(inv, K.INVEST, Decimal("1000000"), timezone.localdate(), self.treasurer)
        self.login("treasurer")
        html = self.client.get(f"/reports/cycle/{self.cycle.pk}/").content.decode()
        self.assertNotIn("</script><script>alert(1)", html)
        self.assertNotIn("<script>alert(1)", html)
        html = self.client.get("/investments/").content.decode()
        self.assertNotIn("<script>alert(1)", html)
