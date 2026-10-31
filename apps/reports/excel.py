"""
Excel exports (openpyxl).
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0'


def safe(value):
    """
    Text typed by users can begin with = + - @ and would be run as a formula when
    the file is opened in Excel. Prefix such text so it stays text.
    """

    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value

    return value


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)

    return safe(value)


def add_sheet(workbook, title, headers, rows, money_columns=(), first=False):
    """`money_columns` are 0-based column numbers shown as whole shillings."""

    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title[:31]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        sheet.append([_cell(v) for v in row])

    for index in money_columns:
        for cell in sheet.iter_rows(min_row=2, min_col=index + 1, max_col=index + 1):
            cell[0].number_format = MONEY

    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1])) for r in rows if index - 1 < len(r)] or [10])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 48)

    sheet.freeze_panes = "A2"

    return sheet


def to_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def cycle_workbook(report):
    cycle = report["cycle"]
    c = report["contributions"]
    loans = report["loans"]
    inv = report["investments"]
    pos = report["position"]
    exp = report["expenses"]

    wb = Workbook()

    summary_rows = [
        ["Cycle", f"Cycle {cycle.cycle_number} - {cycle.name}"],
        ["Period", f"{cycle.start_date:%d %b %Y} to {cycle.end_date:%d %b %Y}"],
        ["Generated", f"{report['generated_at']:%d %b %Y %H:%M}"],
        [],
        ["CONTRIBUTIONS"],
        ["Eligible members", c["eligible_members"]],
        ["Deposit target", c["expected"]],
        ["Deposits collected", c["collected"]],
        ["Deposits outstanding", c["outstanding"]],
        ["Emergency obligation", c["emergency_expected"]],
        ["Emergency collected", c["emergency_collected"]],
        ["Members who met the target", f"{c['members_met_target']} of {c['members']}"],
        [],
        ["LOANS (disbursed in this cycle)"],
        ["Loans", loans["count"]],
        ["Disbursed", loans["disbursed"]],
        ["Repayments received", loans["repayments_received"]],
        ["Recovered from savings", loans["recovered_from_savings"]],
        ["Interest earned", loans["interest_earned"]],
        ["Outstanding", loans["outstanding"]],
        [],
        ["INVESTMENTS"],
        ["Invested this cycle", inv["invested_in_cycle"]],
        ["Earnings this cycle", inv["income_in_cycle"]],
        ["Worth now", inv["market_value"]],
        ["Profit shared with members", inv["distributed"]],
        [],
        ["EXPENSES"],
        ["Spent this cycle", exp["total"]],
        *[[f"  {row['label']}", row["total"]] for row in exp["by_category"]],
        [],
        ["FINANCIAL POSITION (as of now)"],
        ["Members' deposits", pos["members_deposits"]],
        ["Emergency fund", pos["emergency_fund"]],
        ["Loans receivable", pos["loans_receivable"]],
        ["Investments", pos["investments"]],
        ["Derived cash (includes expenses, excludes membership fees)", pos["derived_cash"]],
    ]
    sheet = add_sheet(wb, "Summary", ["Item", "Value"], summary_rows, first=True)
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        if row[0].value and str(row[0].value).isupper():
            row[0].font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(row[0].value, float):
            row[0].number_format = MONEY
    sheet.column_dimensions["A"].width = 52
    sheet.column_dimensions["B"].width = 32

    add_sheet(
        wb, "Members",
        ["Member", "Name", "Status", "Deposit target", "Deposits paid", "Recovered for loans",
         "Deposits outstanding", "Emergency target", "Emergency paid", "Emergency outstanding", "Target met"],
        [
            [r["member"].member_code, r["member"].full_name, r["status"], r["deposit_target"], r["deposits_paid"],
             r["recovered"], r["deposits_outstanding"], r["emergency_target"], r["emergency_paid"],
             r["emergency_outstanding"], "Yes" if r["target_met"] else "No"]
            for r in report["members"]
        ],
        money_columns=(3, 4, 5, 6, 7, 8, 9),
    )

    add_sheet(
        wb, "Loans",
        ["Loan", "Member", "Disbursed", "Final due", "Principal", "Total to repay", "Paid", "Outstanding",
         "Interest paid", "Status"],
        [
            [r["loan"].reference, r["loan"].member.member_code, r["loan"].disbursed_on, r["loan"].final_due_date,
             r["principal"], r["total_due"], r["paid"], r["outstanding"], r["interest_paid"], r["status"]]
            for r in loans["rows"]
        ],
        money_columns=(4, 5, 6, 7, 8),
    )

    add_sheet(
        wb, "Investments",
        ["Investment", "Type", "Provider", "Invested", "Returned", "Earnings", "Fees", "Worth now", "Total return", "Return %"],
        [
            [p["investment"].name, p["investment"].get_kind_display(), p["investment"].provider, p["invested"],
             p["returned"], p["income"], p["fees"], p["market_value"], p["total_return"], p["return_percent"]]
            for p in inv["positions"]
        ],
        money_columns=(3, 4, 5, 6, 7, 8),
    )

    add_sheet(
        wb, "Expenses",
        ["Date", "Category", "Description", "Paid to", "Method", "Reference", "Amount", "Recorded by"],
        [
            [e.expense_date, e.get_category_display(), e.description, e.paid_to, e.get_payment_method_display(),
             e.reference, e.amount, str(e.recorded_by)]
            for e in exp["rows"]
        ],
        money_columns=(6,),
    )

    return to_bytes(wb)


def contributions_workbook(contributions):
    wb = Workbook()
    add_sheet(
        wb, "Contributions",
        ["Date paid", "Member", "Name", "Cycle", "Counts towards", "Type", "For month", "Amount", "Method",
         "Reference", "Recorded by", "Reversed"],
        [
            [c.payment_date, c.member.member_code, c.member.full_name, f"Cycle {c.cycle.cycle_number}",
             f"Cycle {c.settles_cycle.cycle_number} (arrears)" if c.settles_cycle_id else f"Cycle {c.cycle.cycle_number}",
             c.get_contribution_type_display(), c.period_month.strftime("%b %Y"), c.amount,
             c.get_payment_method_display(), c.reference, str(c.recorded_by),
             c.reversed_at.strftime("%d %b %Y") if c.reversed_at else ""]
            for c in contributions
        ],
        money_columns=(7,), first=True,
    )
    return to_bytes(wb)


def loans_workbook(loans):
    wb = Workbook()
    add_sheet(
        wb, "Loans",
        ["Loan", "Member", "Name", "Disbursed", "Final due", "Principal", "Total to repay", "Paid",
         "Outstanding", "Status"],
        [
            [l.reference, l.member.member_code, l.member.full_name, l.disbursed_on, l.final_due_date, l.principal,
             l.total_due, l.total_paid, l.outstanding, l.display_status()]
            for l in loans
        ],
        money_columns=(5, 6, 7, 8), first=True,
    )
    rows = []
    for l in loans:
        for r in l.repayments.all():
            rows.append([
                l.reference, l.member.member_code, r.payment_date, r.get_kind_display(), r.amount,
                r.get_payment_method_display() if r.payment_method else "", r.reference,
                str(r.recorded_by) if r.recorded_by else "System (automatic)",
                f"Reversed: {r.reversal_reason}" if r.reversed_at else "",
            ])
    add_sheet(
        wb, "Repayments",
        ["Loan", "Member", "Date", "What", "Amount", "Method", "Reference", "Recorded by", "Reversed"],
        rows, money_columns=(4,),
    )
    return to_bytes(wb)


def investments_workbook(book, transactions):
    wb = Workbook()
    add_sheet(
        wb, "Portfolio",
        ["Investment", "Type", "Provider", "Status", "Invested", "Returned", "Earnings", "Fees", "Worth now", "Total return"],
        [
            [p["investment"].name, p["investment"].get_kind_display(), p["investment"].provider,
             p["investment"].get_status_display(), p["invested"], p["returned"], p["income"], p["fees"],
             p["market_value"], p["total_return"]]
            for p in book["positions"]
        ],
        money_columns=(4, 5, 6, 7, 8, 9), first=True,
    )
    add_sheet(
        wb, "Transactions",
        ["Date", "Investment", "Type", "Amount", "Reference", "Notes", "Recorded by", "Reversed"],
        [
            [t.transaction_date, t.investment.name, t.get_kind_display(), t.amount, t.reference, t.notes,
             str(t.recorded_by), f"Reversed: {t.reversal_reason}" if t.reversed_at else ""]
            for t in transactions
        ],
        money_columns=(3,),
    )
    return to_bytes(wb)


def statement_workbook(statement):
    member, cycle = statement["member"], statement["cycle"]
    wb = Workbook()
    add_sheet(
        wb, "Summary", ["Item", "Value"],
        [
            ["Member", f"{member.member_code} - {member.full_name}"],
            ["Cycle", f"Cycle {cycle.cycle_number} - {cycle.name}"],
            ["Deposit target", statement["deposit_target"]],
            ["Deposits paid", statement["deposits_paid"]],
            ["Recovered for loans", statement["deposits_offsets"]],
            ["Deposits outstanding", statement["deposits_outstanding"]],
            ["Emergency target", statement["emergency_target"]],
            ["Emergency paid", statement["emergency_paid"]],
            ["Emergency outstanding", statement["emergency_outstanding"]],
            ["Available savings", statement["available_savings"]],
            ["Loan outstanding", statement["loan_outstanding"]],
        ],
        first=True,
    )
    add_sheet(
        wb, "Ledger",
        ["Date", "Description", "Method", "Reference", "Deposit", "Emergency", "Deposits to date", "Note"],
        [
            [r["date"], r["description"], r["method"], r["reference"],
             r["amount"] if r["is_deposit"] else "", "" if r["is_deposit"] else r["amount"],
             r["running_deposits"], r["note"]]
            for r in statement["rows"]
        ],
        money_columns=(4, 5, 6),
    )
    return to_bytes(wb)
