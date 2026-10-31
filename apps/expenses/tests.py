from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import PermissionDenied
from django.test import TestCase
from django.utils import timezone

from apps.accounts.models import User
from apps.auditlog.models import AuditLog
from apps.cycles import services as cycle_services
from apps.cycles.models import Cycle, _add_months

from . import services
from .models import Expense


class ExpenseCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        mk = lambda name, role: User.objects.create_user(username=name, password="pw", role=role)
        cls.admin = mk("admin", User.Role.ADMINISTRATOR)
        cls.treasurer = mk("treasurer", User.Role.TREASURER)
        cls.chair = mk("chair", User.Role.CHAIRPERSON)
        cls.secretary = mk("secretary", User.Role.SECRETARY)
        cls.plain = mk("plain", User.Role.MEMBER)
        first = timezone.localdate().replace(day=1)
        cls.cycle = cycle_services.create_cycle(
            Cycle(cycle_number=1, start_date=_add_months(first, -2), duration_months=6), actor=cls.admin)
        cycle_services.activate_cycle(cls.cycle, actor=cls.admin)

    def make(self, amount="25000", category="TRANSPORT", actor=None, day=None):
        return services.record_expense(
            Expense(cycle=self.cycle, category=category, description="Bus fare to the bank", amount=Decimal(amount),
                    expense_date=day or timezone.localdate(), payment_method="CASH"),
            actor or self.treasurer)


class RecordingTests(ExpenseCase):
    def test_record_sets_the_recorder_and_is_audited(self):
        e = self.make()
        self.assertEqual(e.recorded_by, self.treasurer)
        self.assertTrue(AuditLog.objects.filter(model_name="Expense", object_id=str(e.pk), action="CREATE").exists())

    def test_only_treasurer_and_administrator_record(self):
        self.make(actor=self.admin)
        for user in (self.chair, self.secretary, self.plain):
            with self.assertRaises(PermissionDenied):
                self.make(actor=user)

    def test_bad_amounts_and_future_dates_are_refused(self):
        with self.assertRaises(services.ExpenseError):
            self.make(amount="0")
        with self.assertRaises(services.ExpenseError):
            self.make(day=timezone.localdate() + timedelta(days=1))

    def test_edit_needs_a_reason_and_keeps_the_old_values(self):
        e = self.make()
        old = services._snapshot(e)
        e.amount = Decimal("30000")
        with self.assertRaises(services.ExpenseError):
            services.update_expense(e, old, self.treasurer, " ")
        services.update_expense(e, old, self.treasurer, "typo")
        entry = AuditLog.objects.filter(model_name="Expense", action="UPDATE").latest("id")
        self.assertEqual(entry.old_value["amount"], "25000")
        self.assertEqual(entry.new_value["reason"], "typo")

    def test_delete_needs_a_reason_and_keeps_a_copy(self):
        e = self.make()
        with self.assertRaises(services.ExpenseError):
            services.delete_expense(e, self.treasurer, "")
        with self.assertRaises(PermissionDenied):
            services.delete_expense(e, self.secretary, "x")
        services.delete_expense(e, self.treasurer, "entered twice")
        self.assertFalse(Expense.objects.filter(pk=e.pk).exists())
        self.assertTrue(AuditLog.objects.filter(new_value__event="EXPENSE_DELETED").exists())

    def test_a_closed_cycle_is_read_only(self):
        e = self.make()
        cycle_services.close_cycle(self.cycle, actor=self.admin)
        with self.assertRaises(services.ExpenseError):
            services.delete_expense(e, self.admin, "late")
        with self.assertRaises(services.ExpenseError):
            services.update_expense(e, services._snapshot(e), self.admin, "late")


class SummaryAndReportTests(ExpenseCase):
    def test_summary_totals_by_category(self):
        self.make("10000", "TRANSPORT")
        self.make("5000", "TRANSPORT")
        self.make("40000", "MEETING")
        s = services.expense_summary(self.cycle)
        self.assertEqual(s["total"], Decimal("55000"))
        self.assertEqual(s["count"], 3)
        self.assertEqual([r["category"] for r in s["by_category"]], ["MEETING", "TRANSPORT"])

    def test_cycle_report_excel_and_pdf_include_expenses(self):
        from openpyxl import load_workbook

        from apps.reports import excel, pdf
        from apps.reports.services import cycle_report

        self.make("12345", "STATIONERY")
        report = cycle_report(self.cycle)
        self.assertEqual(report["expenses"]["total"], Decimal("12345"))
        wb = load_workbook(BytesIO(excel.cycle_workbook(report)))
        self.assertIn("Expenses", wb.sheetnames)
        self.assertEqual(wb["Expenses"]["C2"].value, "Bus fare to the bank")
        self.assertTrue(pdf.render_cycle_pdf(report).startswith(b"%PDF"))

    def test_expenses_reduce_the_derived_cash(self):
        from apps.reports.services import financial_position

        before = financial_position()["derived_cash"]
        self.make("20000")
        self.assertEqual(financial_position()["derived_cash"], before - Decimal("20000"))


class ViewTests(ExpenseCase):
    def test_officers_see_the_list_and_the_nav_link(self):
        self.make("15000")
        self.client.force_login(self.chair)
        page = self.client.get("/expenses/")
        self.assertContains(page, "Total spent this cycle")
        self.assertContains(page, "Bus fare to the bank")
        self.assertContains(page, "By category")

    def test_members_and_secretary_cannot_open_it(self):
        for user in (self.plain, self.secretary):
            self.client.force_login(user)
            self.assertEqual(self.client.get("/expenses/").status_code, 403)

    def test_treasurer_records_edits_and_deletes_over_http(self):
        self.client.force_login(self.treasurer)
        r = self.client.post("/expenses/new/", {
            "category": "MEETING", "description": "Tea for the AGM", "amount": "30000",
            "expense_date": timezone.localdate().isoformat(), "payment_method": "CASH"})
        self.assertEqual(r.status_code, 302)
        e = Expense.objects.get(description="Tea for the AGM")
        self.assertEqual(self.client.get(f"/expenses/{e.pk}/edit/").status_code, 200)
        self.client.post(f"/expenses/{e.pk}/edit/", {
            "category": "MEETING", "description": "Tea for the AGM", "amount": "35000",
            "expense_date": timezone.localdate().isoformat(), "payment_method": "CASH", "reason": "receipt says 35,000"})
        e.refresh_from_db()
        self.assertEqual(e.amount, Decimal("35000"))
        self.client.post(f"/expenses/{e.pk}/delete/", {"reason": "duplicate"})
        self.assertFalse(Expense.objects.filter(pk=e.pk).exists())

    def test_the_report_page_has_the_expenses_section(self):
        self.make("9000")
        self.client.force_login(self.admin)
        page = self.client.get(f"/reports/cycle/{self.cycle.pk}/")
        self.assertContains(page, "4. Expenses")
